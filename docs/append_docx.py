#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Дописывает в отчёт «3. Проектирование архитектуры системы.docx» диаграммы
классов (PNG из klass/) с подписями «Рис. N» и гиперссылками на исходники."""
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

DOCX = "3. Проектирование архитектуры системы.docx"
IMG_DIR = "klass"


def add_hyperlink(paragraph, url, text, color="0563C1", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rPr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_picture_centered(doc, path, width_in=6.0):
    doc.add_picture(path, width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(11)


def add_source(doc, path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Источник: ")
    r.font.size = Pt(10)
    add_hyperlink(p, path, path)


figures = [
    ("klass-1.png", "Диаграмма классов фронтенда (общая)", "klass/klass.md"),
    ("klass_frontend_1_core-1.png", "Слой 1 — Ядро и конфигурация", "klass/klass_frontend_1_core.md"),
    ("klass_frontend_2_state-1.png", "Слой 2 — Состояние и случайность", "klass/klass_frontend_2_state.md"),
    ("klass_frontend_3_render-1.png", "Слой 3 — Отрисовка и рендеринг", "klass/klass_frontend_3_render.md"),
    ("klass_frontend_4_tools-1.png", "Слой 4 — Инструменты и генератор", "klass/klass_frontend_4_tools.md"),
    ("klass_frontend_5_io_cloud-1.png", "Слой 5 — IO, облако и публикация", "klass/klass_frontend_5_io_cloud.md"),
    ("klass_frontend_6_ui-1.png", "Слой 6 — UI, авторизация и галерея", "klass/klass_frontend_6_ui.md"),
    ("klass_backend-1.png", "Диаграмма классов бэкенда (общая)", "klass/klass_backend.md"),
    ("klass_backend_1_app-1.png", "Слой 1 — Каркас приложения", "klass/klass_backend_1_app.md"),
    ("klass_backend_2_models-1.png", "Слой 2 — Модели данных", "klass/klass_backend_2_models.md"),
    ("klass_backend_3_repos-1.png", "Слой 3 — Репозитории", "klass/klass_backend_3_repos.md"),
    ("klass_backend_4_service-1.png", "Слой 4 — Сервис галереи", "klass/klass_backend_4_service.md"),
    ("klass_backend_5_api-1.png", "Слой 5 — API-маршруты", "klass/klass_backend_5_api.md"),
]


def main():
    doc = Document(DOCX)

    doc.add_page_break()
    doc.add_heading("Приложение А — Диаграммы классов", level=1)
    doc.add_paragraph(
        "В приложении приведены диаграммы классов системы «Dungeon Map Builder»: "
        "общая диаграмма классов фронтенда и её разбиение на слои (поддиаграммы), "
        "а также диаграмма классов бэкенда (Flask) с разбиением на слои. "
        "Каждая диаграмма сопровождается ссылкой на исходный файл в формате Mermaid."
    )

    doc.add_heading("А.1. Фронтенд", level=2)
    for i, (png, caption, src) in enumerate(figures, start=1):
        if i == 8:
            doc.add_heading("А.2. Бэкенд", level=2)
        png_path = os.path.join(IMG_DIR, png)
        if not os.path.exists(png_path):
            print("MISSING:", png_path)
            continue
        add_picture_centered(doc, png_path)
        add_caption(doc, f"Рис. {i} — {caption}")
        add_source(doc, src)

    doc.save(DOCX)
    print("OK ->", DOCX)


if __name__ == "__main__":
    main()
